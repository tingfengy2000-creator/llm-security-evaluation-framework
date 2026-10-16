"""Owner-authorized, exact-text D1 Human comparison; no correction or GT.

Inputs are SHA-bound raw returns, own visible issued books and own Phase1.
No construction mapping/roles, Expected, reviewer answers or live source reads.
"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any

from scripts.research.validate_core144_d1_human_phase2_returns import (
    DOC, FIELDS, PHASE1, SHEET, new_json, ordered_sha, readonly_differences,
    row_checks, sha, validation_signature, values, zip_check,
)

OLD = Path("experiments/core144_d1_human_phase2_raw_lock_20261016")
OUT = Path("experiments/core144_d1_human_ab_comparison_20261016")
INDEX_SHA = "6ee7a95a0bf200255c9afc3a38ddb08c82f05ff0f68fa7c6da4113c55fbafb36"
DIFF_SHA = "1e5a9747598e9701d0db9c27ae1e85a1ae102f84f306fff1bebbe1758c420ec8"
P1_FIELDS = ("text_naturalness", "local_internal_conflict", "self_containment", "ambiguous_referent", "meta_or_template_language")
SEMANTIC = tuple(f for f in FIELDS[:8] if f != "evidence_selection")


def load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def exact_pairs(a: list[dict[str, Any]], b: list[dict[str, Any]], count: int = 144) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    if len(a) != count or len(b) != count:
        raise ValueError("Pairing record count mismatch")
    left = [r["candidate_text"] for r in a]
    right = [r["candidate_text"] for r in b]
    if any(not isinstance(t, str) or not t for t in left + right):
        raise ValueError("Missing Candidate text")
    if len(set(left)) != count or len(set(right)) != count or set(left) != set(right):
        raise ValueError("Exact unique Candidate-text bijection required; never trim")
    lookup = {r["candidate_text"]: r for r in b}
    return [(r, lookup[r["candidate_text"]]) for r in a]


def exception_applies(raw_sha: str, approved_raw: str, issued_sha: str, approved_issued: str,
                      actual: list[dict[str, Any]], frozen: list[dict[str, Any]]) -> bool:
    return (raw_sha == approved_raw and issued_sha == approved_issued and len(actual) == 123
            and actual == frozen
            and all(d["kind"] == "BOUNDARY_WHITESPACE_ONLY" and isinstance(d["before"], str)
                    and isinstance(d["after"], str) and d["before"].strip() == d["after"].strip()
                    for d in actual))


def agreement(a: list[str], b: list[str]) -> dict[str, Any]:
    if not a or len(a) != len(b):
        raise ValueError("Nonempty paired observations required")
    n = len(a)
    ca, cb = Counter(a), Counter(b)
    agree_count = sum(x == y for x, y in zip(a, b))
    expected_numerator = sum(ca[k] * cb[k] for k in ca.keys() | cb.keys())
    denominator = n * n - expected_numerator
    cells = Counter(zip(a, b))
    return {"agree": agree_count, "total": n, "disagree": n - agree_count,
            "agreement_rate": agree_count / n,
            "cohen_kappa": (agree_count * n - expected_numerator) / denominator if denominator else None,
            "kappa_status": "COMPUTED" if denominator else "NOT_ESTIMABLE_CONSTANT_MARGINALS",
            "a_marginals": dict(sorted(ca.items())), "b_marginals": dict(sorted(cb.items())),
            "confusion_counts": [{"a": x, "b": y, "count": count} for (x, y), count in sorted(cells.items())]}


def partitions(rows: list[dict[str, Any]]) -> dict[str, int]:
    return {"semantic_rows": sum(bool(r["semantic_differences"]) for r in rows),
            "process_only_rows": sum(not r["semantic_differences"] and r["process_difference"] for r in rows),
            "all_categorical_equal_rows": sum(not r["semantic_differences"] and not r["process_difference"] for r in rows)}


def verify_old_index() -> dict[str, Any]:
    path = OLD / "EVIDENCE_INDEX_V1.json"
    if sha(path.read_bytes()) != INDEX_SHA:
        raise ValueError("Old sealed index identity mismatch")
    index = load(path)
    for item in index["files"]:
        source = (OLD / item["path"]).resolve()
        if not source.is_relative_to(OLD.resolve()):
            raise ValueError("Unsafe historical index path")
        if source.stat().st_size != item["bytes"] or sha(source.read_bytes()) != item["sha256"]:
            raise ValueError(f"Historical evidence changed: {item['path']}")
    return {"index_sha256": INDEX_SHA, "verified_files": len(index["files"]), "unchanged": True}


def read_person(lock: dict[str, Any], approval: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    person = lock["annotator"]
    suffix = person.replace("-", "_")
    issued_meta = load(DOC / f"PAPER1_CORE144_D1_{suffix}_PHASE2_MANIFEST_V1.json")
    raw_path, issued_path = Path(lock["locked_path"]), Path(issued_meta["destination"])
    p1_path, p1_sha = PHASE1[person]
    inputs = [(raw_path, lock["sha256"]), (issued_path, issued_meta["sha256"]), (p1_path, p1_sha)]
    for path, wanted in inputs:
        if sha(path.read_bytes()) != wanted or not zip_check(path)["pass"]:
            raise ValueError(f"Own input identity/safety failure: {person}, {path.name}")
    raw, issued, p1 = [load_workbook(path, data_only=False) for path, _ in inputs]
    sheet, base, first = raw[SHEET], issued[SHEET], p1[SHEET]
    if raw.sheetnames != issued.sheetnames or (sheet.max_row, sheet.max_column) != (145, 20):
        raise ValueError("Workbook shape drift")
    if [c.value for c in sheet[1]] != [c.value for c in base[1]]:
        raise ValueError("Header drift")
    ids = [sheet.cell(r, 1).value for r in range(2, 146)]
    texts = [sheet.cell(r, 2).value for r in range(2, 146)]
    if len(set(ids)) != 144 or len(set(texts)) != 144:
        raise ValueError("Own duplicate identity")
    if ids != [base.cell(r, 1).value for r in range(2, 146)] or texts != [base.cell(r, 2).value for r in range(2, 146)]:
        raise ValueError("Own issued ID/order/Candidate drift")
    if ids != [first.cell(r, 1).value for r in range(2, 146)] or texts != [first.cell(r, 2).value for r in range(2, 146)]:
        raise ValueError("Own Phase1 identity drift")
    if ordered_sha(ids) != issued_meta["ordered_ids_sha256"] or ordered_sha(texts) != issued_meta["candidate_texts_in_order_sha256"]:
        raise ValueError("Own frozen order SHA mismatch")
    differences = readonly_differences(values(issued), values(raw))
    for d in differences:
        if d["sheet"] == SHEET and 2 <= d["row"] <= 145:
            d["blind_id"] = sheet.cell(d["row"], 1).value
    frozen = load(OLD / f"{suffix}_READONLY_CELL_DIFFERENCES_V1.json")
    if differences != frozen:
        raise ValueError("New readonly deviation beyond frozen audit")
    if person == "HUMAN-A01":
        if sha((OLD / "HUMAN_A01_READONLY_CELL_DIFFERENCES_V1.json").read_bytes()) != DIFF_SHA:
            raise ValueError("Approved cell-list SHA mismatch")
        allowed = exception_applies(lock["sha256"], approval["raw_sha256"], issued_meta["sha256"], approval["issued_sha256"], differences, frozen)
    else:
        allowed = all(d["kind"] == "BLANK_STORAGE_EQUIVALENT" for d in differences)
    if not allowed:
        raise ValueError("Presentation gate not authorized for these exact bytes/cells")
    if validation_signature(sheet) != validation_signature(base):
        raise ValueError("Dropdown contract drift")
    def native_signature(book: Any) -> Any:
        return ({(s.title, c.coordinate): c.hyperlink.target for s in book for c in s._cells.values() if c.hyperlink},
                {(s.title, c.coordinate): c.value for s in book for c in s._cells.values() if c.data_type == "f"},
                [(s.title, s.sheet_state) for s in book],
                [(r, d.hidden) for r, d in book[SHEET].row_dimensions.items() if d.hidden],
                [(c, d.hidden) for c, d in book[SHEET].column_dimensions.items() if d.hidden])
    if native_signature(raw) != native_signature(issued):
        raise ValueError("Native link/formula/visibility drift")
    result = []
    legal = 0
    for r in range(2, 146):
        answers = {f: sheet.cell(r, c).value for c, f in enumerate(FIELDS, 11)}
        phase1 = {f: first.cell(r, c).value for c, f in enumerate(P1_FIELDS, 3)}
        phase1["issue_note"] = first.cell(r, 8).value
        available = {tag for tag, col in (("E1", 6), ("E2", 10)) if base.cell(r, col).value not in (None, "")}
        checks = row_checks(answers, phase1["local_internal_conflict"], available)
        if checks["errors"] or checks["consistency_flags"]:
            raise ValueError("New own structural/crossfield blocker")
        legal += checks["legal_categorical_cells"]
        result.append({"blind_id": ids[r-2], "excel_row": r, "candidate_text": texts[r-2],
                       "candidate_text_sha256": sha(texts[r-2].encode()), "phase1": phase1, "phase2": answers,
                       "own_review_leads": checks["review_leads"],
                       "evidence_bundle": {tag: {"title": base.cell(r, col).value, "excerpt": base.cell(r, col+1).value,
                                                 "official_url_cell": base.cell(r, col+2).value, "snapshot_ref": base.cell(r, col+3).value}
                                           for tag, col in (("E1", 3), ("E2", 7))}})
    for path, wanted in inputs:
        if sha(path.read_bytes()) != wanted:
            raise ValueError("Input changed during read")
    for book in (raw, issued, p1):
        book.close()
    return result, {"annotator": person, "rows": len(result), "legal_categorical_cells": legal,
                    "raw_sha256": lock["sha256"], "issued_sha256": issued_meta["sha256"], "phase1_sha256": p1_sha,
                    "strict_literal_nonblank_parity": person != "HUMAN-A01",
                    "approved_presentation_exception_applied": person == "HUMAN-A01",
                    "effective_receipt_gate_pass": allowed, "native_parity": True,
                    "raw_and_issued_unchanged_after_read": True}


def compare(a: list[dict[str, Any]], b: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    pairs = exact_pairs(a, b)
    rows = []
    for index, (left, right) in enumerate(pairs, 1):
        if left["evidence_bundle"] != right["evidence_bundle"]:
            raise ValueError("Assigned A/B E1/E2 bundle mismatch; local aliases not comparable")
        differences = [f for f in SEMANTIC if left["phase2"][f] != right["phase2"][f]]
        first = [f for f in P1_FIELDS if left["phase1"][f] != right["phase1"][f]]
        process = left["phase2"]["evidence_selection"] != right["phase2"]["evidence_selection"]
        leads = [{"person": who, "code": c} for who, rec in (("HUMAN-A01", left), ("HUMAN-B01", right)) for c in rec["own_review_leads"]]
        rows.append({"comparison_id": f"D1HC-{index:04d}", "candidate_text_sha256": left["candidate_text_sha256"],
                     "candidate_text": left["candidate_text"], "a": left, "b": right,
                     "phase1_differences": first, "semantic_differences": differences, "process_difference": process,
                     "review_leads": leads, "needs_owner_review": bool(first or differences or leads),
                     "linked_local_minimum_difference": "local_internal_conflict" in first and "minimum_external_evidence_needed" in differences,
                     "proposed_adjudicated_values": None, "effective_adjudicated_values": None})
    phase2 = {f: agreement([r["phase2"][f] for r in a], [y["phase2"][f] for _, y in pairs]) for f in FIELDS[:8]}
    phase1 = {f: agreement([r["phase1"][f] for r in a], [y["phase1"][f] for _, y in pairs]) for f in P1_FIELDS}
    semantic_agree = sum(phase2[f]["agree"] for f in SEMANTIC)
    stats = {"paired_rows": len(rows), "bijection": True, "evidence_bundle_parity": True,
             "phase2_fields": phase2, "phase1_fields": phase1, **partitions(rows),
             "phase2_semantic_cell_agreement": {"agree": semantic_agree, "total": 1008, "rate": semantic_agree / 1008},
             "phase2_all_seven_equal_rows": sum(not r["semantic_differences"] for r in rows),
             "phase1_disagreement_rows": sum(bool(r["phase1_differences"]) for r in rows),
             "quality_lead_rows": sum(bool(r["review_leads"]) for r in rows),
             "owner_review_union_rows": sum(r["needs_owner_review"] for r in rows),
             "linked_local_minimum_difference_rows": sum(r["linked_local_minimum_difference"] for r in rows),
             "phase2_reason_raw_equal_rows": sum(x["phase2"]["phase2_reason"] == y["phase2"]["phase2_reason"] for x, y in pairs),
             "reason_equality_is_not_acceptance_gate": True,
             "adjudication_or_gt_created": False, "expected_mapping_or_construction_roles_loaded": False,
             "interpretation": "Descriptive Human agreement, not annotation accuracy or dataset acceptance"}
    return rows, stats


def main() -> None:
    approval = load(OUT / "OWNER_APPROVAL_RECEIPT_V1.json")
    if (approval["owner_text"] != "批准仅本次123格呈现等价，原件不改，继续A/B比较。"
            or approval["person"] != "HUMAN-A01" or approval["approved_cell_count"] != 123
            or not approval["ab_comparison_authorized"] or approval["gt_authorized"]
            or approval["readonly_difference_list_sha256"] != DIFF_SHA):
        raise ValueError("Exact Owner scope required")
    historical = verify_old_index()
    locks = load(OLD / "DUAL_PHASE2_RAW_LOCK_MANIFEST_V1.json")
    if [r["annotator"] for r in locks] != ["HUMAN-A01", "HUMAN-B01"]:
        raise ValueError("Exact annotator cohort required")
    a, av = read_person(locks[0], approval)
    b, bv = read_person(locks[1], approval)
    overlay = {**approval, "historical_index_verification": historical,
               "strict_literal_parity_remains_false": True, "effective_presentation_gate_pass": True,
               "approved_cells": load(OLD / "HUMAN_A01_READONLY_CELL_DIFFERENCES_V1.json"),
               "raw_rewritten": False, "canonical_rules_changed": False}
    new_json(OUT / "A01_PRESENTATION_EQUIVALENCE_OVERLAY_V1.json", overlay)
    new_json(OUT / "INPUT_PREFLIGHT_V1.json", {"historical": historical, "a": av, "b": bv})
    rows, stats = compare(a, b)
    new_json(OUT / "AB_PAIRED_COMPARISON_V1.json", rows)
    new_json(OUT / "AB_FIELD_AGREEMENT_V1.json", stats)
    new_json(OUT / "OWNER_REVIEW_QUEUE_V1.json", [r for r in rows if r["needs_owner_review"]])
    new_json(OUT / "PROCESS_ONLY_DIFFERENCES_V1.json", [r for r in rows if not r["semantic_differences"] and r["process_difference"]])
    if not verify_old_index()["unchanged"]:
        raise ValueError("Historical evidence changed during comparison")
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
