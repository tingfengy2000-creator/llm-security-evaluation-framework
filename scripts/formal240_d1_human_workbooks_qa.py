"""Read-only QA for the four D1 human workbooks; never edits their bytes."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]


ROOT = Path(__file__).resolve().parents[1]
PHASE1_SHEETS = [
    "【开始前必读】", "【标注表】", "【字段说明】", "【高危歧义与易错点】",
    "【判断流程】", "【虚构示例】", "【提交前自检】",
]
PHASE2_SHEETS = [
    "【开始前必读】", "【标注表】", "【字段说明】", "【高危歧义与易错点】",
    "【版本与历史判断】", "【权威与机关判断】", "【证据判断流程】",
    "【虚构示例】", "【冻结证据全文】", "【提交前自检】",
]
ENUMS_1 = {
    "C": {"NATURAL", "MINOR_ISSUE", "UNNATURAL"},
    "D": {"YES", "NO", "UNCERTAIN"},
    "E": {"PASS", "FLAG", "UNCERTAIN"},
    "F": {"YES", "NO", "UNCERTAIN"},
    "G": {"YES", "NO", "UNCERTAIN"},
}
ENUMS_2 = {
    "K": {"CURRENTLY_CONSISTENT", "LEGITIMATE_VERSION_OR_HISTORY", "FACTUAL_CONFLICT", "INSUFFICIENT_EVIDENCE"},
    "L": {"NOT_PRESENT", "PRESENT_CORRECT", "PRESENT_INCORRECT", "PRESENT_EVIDENCE_INSUFFICIENT"},
    "M": {"NOT_PRESENT", "PRESENT_CORRECT", "PRESENT_INCORRECT", "PRESENT_EVIDENCE_INSUFFICIENT"},
    "N": {"ZERO_EXTERNAL_EVIDENCE_REQUIRED", "ONE_OFFICIAL_EVIDENCE", "MULTI_EVIDENCE_OR_VERSION_CHAIN", "NOT_APPLICABLE"},
    "O": {"NONE", "E1", "E2", "E1+E2"},
    "P": {"NONE", "SOURCE_UNREACHABLE", "SOURCE_CONFLICT", "EVIDENCE_MISSING", "LATE_DISCOVERED_CANDIDATE_DEFECT", "OTHER"},
    "Q": {"YES", "NO", "UNCERTAIN"},
    "R": {"SUFFICIENT", "INSUFFICIENT", "UNCERTAIN"},
}


def check(human: str, stage: str, payload_dir: Path, book_dir: Path) -> dict[str, object]:
    payload = json.loads((payload_dir / f"PAPER1_CORE144_D1_HUMAN_{human}_{stage}_PAYLOAD_V1.json").read_text(encoding="utf-8"))
    records = payload["records"]
    assert len(records) == 144
    path = book_dir / f"PAPER1_CORE144_D1_HUMAN_{human}_{stage}_V1.xlsx"
    blob = path.read_bytes()
    book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    expected_sheets = PHASE1_SHEETS if stage == "PHASE1" else PHASE2_SHEETS
    assert book.sheetnames == expected_sheets, book.sheetnames
    assert all(book[s].sheet_state == "visible" for s in expected_sheets)
    sheet = book["【标注表】"]
    assert sheet.max_row == 145 and sheet.max_column == (8 if stage == "PHASE1" else 20)
    assert sheet.freeze_panes == "C2"
    assert len(sheet.tables) == 1 and sheet.auto_filter is not None
    enums = ENUMS_1 if stage == "PHASE1" else ENUMS_2
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == len(enums)
    for col, allowed in enums.items():
        matches = [v for v in validations if str(v.sqref) == f"{col}2:{col}145"]
        assert len(matches) == 1, (human, stage, col)
        formula = matches[0].formula1 or ""
        assert matches[0].type == "list" and all(x in formula for x in allowed)
    ids = [r["blind_id"] for r in records]
    assert len(set(ids)) == 144
    assert [sheet[f"A{i}"].value for i in range(2, 146)] == ids
    assert [sheet[f"B{i}"].value for i in range(2, 146)] == [r["candidate_text"] for r in records]
    for row in range(2, 146):
        for column_index in range(3 if stage == "PHASE1" else 11, sheet.max_column + 1):
            assert sheet.cell(row, column_index).value is None, (human, stage, row, column_index)
    links = 0
    if stage == "PHASE1":
        for tab in book:
            for cells in tab:
                for cell in cells:
                    if isinstance(cell.value, str):
                        assert "https://" not in cell.value
                        assert not re.search(r"\bD1BR-[A-F0-9]{12}\b", cell.value)
    else:
        full = book["【冻结证据全文】"]
        assert full.max_row == 145 and full.max_column == 5
        for i, record in enumerate(records, start=2):
            for col, e in (("E", "E1"), ("I", "E2")):
                url = record[e]["official_url"]
                value = sheet[f"{col}{i}"].value
                if url:
                    assert value == f'=HYPERLINK("{url}","{url}")', (i, col)
                    links += 1
                else:
                    assert value in (None, ""), (i, col)
            assert full[f"A{i}"].value == record["blind_id"]
            assert full[f"C{i}"].value == record["E1"]["excerpt"]
            assert (full[f"E{i}"].value or "") == record["E2"]["excerpt"]
            assert sheet[f"C{i}"].value == record["E1"]["title"]
            assert (sheet[f"G{i}"].value or "") == record["E2"]["title"]
            assert record["E1"]["snapshot_ref"] in sheet[f"F{i}"].value
            if record["E2"]["official_url"]:
                assert record["E2"]["snapshot_ref"] in sheet[f"J{i}"].value
    assert not any("Expected" in name or "Mapping" in name for name in book.sheetnames)
    for tab in book:
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    assert not re.search(r"\bD1BR-[A-F0-9]{12}\b", cell.value)
                    assert "source_blind_id" not in cell.value
                    assert "target_stealth" not in cell.value
                    assert "ground_truth" not in cell.value.lower()
    return {
        "file": path.name,
        "sha256": hashlib.sha256(blob).hexdigest(),
        "bytes": len(blob),
        "sheet_count": len(book.sheetnames),
        "candidate_rows": len(records),
        "enum_dropdowns": len(validations),
        "clickable_url_formulas": links,
        "ids_sha256": hashlib.sha256("\n".join(ids).encode()).hexdigest(),
        "candidate_texts_sha256": hashlib.sha256("\n".join(r["candidate_text"] for r in records).encode()).hexdigest(),
        "xlsx_reimport": "PASS",
    }


def main() -> None:
    payload_dir = ROOT / sys.argv[1]
    book_dir = ROOT / sys.argv[2]
    results = [check(h, s, payload_dir, book_dir) for h in ("A01", "B01") for s in ("PHASE1", "PHASE2")]
    for a, b in ((results[0], results[1]), (results[2], results[3])):
        assert a["ids_sha256"] == b["ids_sha256"]
        assert a["candidate_texts_sha256"] == b["candidate_texts_sha256"]
    assert results[0]["ids_sha256"] != results[2]["ids_sha256"]
    a_records = json.loads((payload_dir / "PAPER1_CORE144_D1_HUMAN_A01_PHASE1_PAYLOAD_V1.json").read_text(encoding="utf-8"))["records"]
    b_records = json.loads((payload_dir / "PAPER1_CORE144_D1_HUMAN_B01_PHASE1_PAYLOAD_V1.json").read_text(encoding="utf-8"))["records"]
    assert sorted(r["candidate_text"] for r in a_records) == sorted(r["candidate_text"] for r in b_records)
    assert set(json.loads((payload_dir / "PAPER1_CORE144_D1_HUMAN_A01_PHASE1_PAYLOAD_V1.json").read_text(encoding="utf-8"))["records"][i]["blind_id"] for i in range(144)).isdisjoint(
        json.loads((payload_dir / "PAPER1_CORE144_D1_HUMAN_B01_PHASE1_PAYLOAD_V1.json").read_text(encoding="utf-8"))["records"][i]["blind_id"] for i in range(144)
    )
    print(json.dumps({"status": "PASS", "results": results}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
