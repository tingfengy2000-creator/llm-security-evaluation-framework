"""Read-only, per-annotator D1 Phase2 receipt validation; never join A/B.

Only own issued workbook and own locked Phase1 are read. No construction
mapping, Expected, GT, reviewer answers, source URL or workbook save is used.
Boundary whitespace equivalence is a diagnostic, NOT a parity waiver.
"""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

FIELDS = (
    "overall_fact_status", "version_claim_status", "authority_claim_status",
    "minimum_external_evidence_needed", "evidence_selection", "phase2_issue",
    "possible_accidental_secondary_error", "evidence_sufficiency",
    "phase2_reason", "reviewer_note",
)
ENUMS = {
    FIELDS[0]: {"CURRENTLY_CONSISTENT", "LEGITIMATE_VERSION_OR_HISTORY", "FACTUAL_CONFLICT", "INSUFFICIENT_EVIDENCE"},
    FIELDS[1]: {"NOT_PRESENT", "PRESENT_CORRECT", "PRESENT_INCORRECT", "PRESENT_EVIDENCE_INSUFFICIENT"},
    FIELDS[2]: {"NOT_PRESENT", "PRESENT_CORRECT", "PRESENT_INCORRECT", "PRESENT_EVIDENCE_INSUFFICIENT"},
    FIELDS[3]: {"ZERO_EXTERNAL_EVIDENCE_REQUIRED", "ONE_OFFICIAL_EVIDENCE", "MULTI_EVIDENCE_OR_VERSION_CHAIN", "NOT_APPLICABLE"},
    FIELDS[4]: {"NONE", "E1", "E2", "E1+E2"},
    FIELDS[5]: {"NONE", "SOURCE_UNREACHABLE", "SOURCE_CONFLICT", "EVIDENCE_MISSING", "LATE_DISCOVERED_CANDIDATE_DEFECT", "OTHER"},
    FIELDS[6]: {"YES", "NO", "UNCERTAIN"},
    FIELDS[7]: {"SUFFICIENT", "INSUFFICIENT", "UNCERTAIN"},
}
SHEET = "【标注表】"
ROOT = Path("experiments/core144_d1_human_phase2_raw_lock_20261016")
DOC = Path("docs/research/stage6_1_hidden_knowledge_poisoning/core144")
PHASE1 = {
    "HUMAN-A01": (Path("E:/LLMGuard-Handoff/paper1_core144_d1_human_a01_phase1_raw_lock_20261002/PAPER1_CORE144_D1_HUMAN_A01_PHASE1_RAW_RETURN_V1.xlsx"), "d752d27451fcbae9ffba82b6bb33445452ef8b51e35f68877793687216f21b28"),
    "HUMAN-B01": (Path("experiments/core144_d1_b01_phase1_raw_lock_20261006/PAPER1_CORE144_D1_HUMAN_B01_PHASE1_RAW_RETURN_V1.xlsx"), "df423484092733066cbed628dde7ced0494dcb8a024c16bd35c0f1d9edfcbc6a"),
}


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def ordered_sha(values: list[str]) -> str:
    return sha("\n".join(values).encode("utf-8"))


def new_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as stream:
        stream.write((json.dumps(obj, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))


def difference_kind(before: Any, after: Any) -> str:
    if before == after:
        return "IDENTICAL"
    if before in (None, "") and after in (None, ""):
        return "BLANK_STORAGE_EQUIVALENT"
    if isinstance(before, str) and isinstance(after, str) and before.strip() == after.strip():
        return "BOUNDARY_WHITESPACE_ONLY"
    return "CONTENT_CHANGE"


def readonly_differences(before: dict[tuple[str, int, int], Any],
                         after: dict[tuple[str, int, int], Any]) -> list[dict[str, Any]]:
    result = []
    for key in sorted(set(before) | set(after)):
        sheet, row, col = key
        if sheet == SHEET and 2 <= row <= 145 and 11 <= col <= 20:
            continue
        old, new = before.get(key), after.get(key)
        kind = difference_kind(old, new)
        if kind == "IDENTICAL":
            continue
        result.append({"sheet": sheet, "row": row, "column": col,
                       "before": old, "after": new, "kind": kind,
                       "before_value_sha256": sha(json.dumps(old, ensure_ascii=False).encode()),
                       "after_value_sha256": sha(json.dumps(new, ensure_ascii=False).encode())})
    return result


def row_checks(row: dict[str, Any], local: str, available: set[str]) -> dict[str, Any]:
    errors, flags, leads = [], [], []
    legal = 0
    for field in FIELDS:
        value = row[field]
        if field == "reviewer_note" and value in (None, ""):
            continue
        if not isinstance(value, str) or value.startswith("="):
            errors.append({"field": field, "code": "NONSTRING_OR_FORMULA"})
        elif field in ENUMS:
            if value not in ENUMS[field]:
                errors.append({"field": field, "code": "ILLEGAL_ENUM", "value": value})
            else:
                legal += 1
        elif field == "phase2_reason" and not value.strip():
            errors.append({"field": field, "code": "EMPTY_REQUIRED_REASON"})
    overall, minimum = row[FIELDS[0]], row[FIELDS[3]]
    if overall != "FACTUAL_CONFLICT" and minimum != "NOT_APPLICABLE":
        flags.append("NONCONFLICT_REQUIRES_NOT_APPLICABLE")
    elif overall == "FACTUAL_CONFLICT":
        if local == "YES" and minimum != "ZERO_EXTERNAL_EVIDENCE_REQUIRED":
            flags.append("OWN_PHASE1_LOCAL_YES_REQUIRES_ZERO")
        elif local != "YES" and minimum == "ZERO_EXTERNAL_EVIDENCE_REQUIRED":
            flags.append("ZERO_WITHOUT_OWN_LOCKED_PHASE1_LOCAL_YES")
        elif minimum == "NOT_APPLICABLE":
            flags.append("CONFLICT_MINIMUM_NOT_APPLICABLE")
    if local == "YES" and overall != "FACTUAL_CONFLICT":
        leads.append("CROSS_PHASE_LOCAL_YES_OVERALL_NOT_CONFLICT_REVIEW")
    selection = row[FIELDS[4]]
    selected = set() if selection == "NONE" or not isinstance(selection, str) else set(selection.split("+"))
    if not selected <= available:
        flags.append("UNSUPPLIED_EVIDENCE_SELECTED")
    if minimum == "MULTI_EVIDENCE_OR_VERSION_CHAIN" and selected != {"E1", "E2"}:
        leads.append("MULTI_MINIMUM_SELECTION_REQUIRES_CONTEXT_REVIEW")
    text = " ".join(str(row[f] or "") for f in ("phase2_reason", "reviewer_note"))
    mentioned = set(re.findall(r"E[12](?!\d)", text))
    if not mentioned <= selected:
        leads.append("REASON_REFERENCE_OUTSIDE_SELECTION_NOT_AUTOMATIC_ERROR")
    for field, ordinary in ((FIELDS[5], "NONE"), (FIELDS[6], "NO"), (FIELDS[7], "SUFFICIENT")):
        if row[field] != ordinary:
            leads.append(f"REVIEW_{field.upper()}")
    return {"errors": errors, "consistency_flags": flags, "review_leads": leads, "legal_categorical_cells": legal}


def strict_receipt_gate(errors: list[Any], differences: list[dict[str, Any]]) -> bool:
    return not errors and not any(d["kind"] != "BLANK_STORAGE_EQUIVALENT" for d in differences)


def comparison_allowed(a: bool, b: bool) -> bool:
    return a and b


def zip_check(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        unsafe = [n for n in names if "vbaProject" in n or n.startswith(("xl/externalLinks/", "xl/embeddings/", "xl/activeX/"))]
        bad = archive.testzip()
    return {"zip_crc_pass": bad is None, "unsafe_parts": unsafe, "pass": bad is None and not unsafe}


def values(book: Any) -> dict[tuple[str, int, int], Any]:
    # Stored cells only; do not allocate beyond original scope during iteration.
    return {(sheet.title, cell.row, cell.column): cell.value
            for sheet in book for cell in sheet._cells.values() if cell.value is not None}


def validation_signature(sheet: Any) -> dict[str, list[str]]:
    result = {}
    for row in range(2, 146):
        for col, field in enumerate(FIELDS[:8], 11):
            cell = sheet.cell(row, col)
            choices = []
            for rule in sheet.data_validations.dataValidation:
                if cell.coordinate in rule and rule.type == "list" and rule.formula1:
                    choices.extend(rule.formula1.strip('"').split(","))
            result[f"{row}:{field}"] = sorted(set(choices))
    return result


def validate_person(lock: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    person = lock["annotator"]
    suffix = person.replace("-", "_")
    manifest_path = DOC / f"PAPER1_CORE144_D1_{suffix}_PHASE2_MANIFEST_V1.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    issued_path = Path(manifest["destination"])
    raw_path = Path(lock["locked_path"])
    p1_path, p1_sha = PHASE1[person]
    inputs = [(issued_path, manifest["sha256"]), (raw_path, lock["sha256"]), (p1_path, p1_sha)]
    for path, expected in inputs:
        if sha(path.read_bytes()) != expected:
            raise ValueError(f"Input SHA mismatch: {person}, {path.name}")
    if issued_path.stat().st_size != manifest["bytes"] or raw_path.stat().st_size != lock["bytes"]:
        raise ValueError("Input byte-count mismatch")
    safety = {"issued": zip_check(issued_path), "raw": zip_check(raw_path), "phase1": zip_check(p1_path)}
    if not all(s["pass"] for s in safety.values()):
        raise ValueError("Unsafe or damaged workbook; do not decode")
    issued = load_workbook(issued_path, data_only=False)
    raw = load_workbook(raw_path, data_only=False)
    p1 = load_workbook(p1_path, data_only=False)
    errors: list[dict[str, Any]] = []
    if raw.sheetnames != issued.sheetnames:
        errors.append({"code": "SHEETS_NOT_EXACT"})
    for s in raw:
        if s.sheet_state != issued[s.title].sheet_state:
            errors.append({"code": "SHEET_VISIBILITY_CHANGE", "sheet": s.title})
    base, returned, first = issued[SHEET], raw[SHEET], p1[SHEET]
    if (returned.max_row, returned.max_column) != (145, 20):
        errors.append({"code": "ANNOTATION_DIMENSIONS_CHANGED"})
    if [c.value for c in returned[1]] != [c.value for c in base[1]]:
        errors.append({"code": "HEADERS_NOT_EXACT"})
    ids = [returned.cell(r, 1).value for r in range(2, 146)]
    texts = [returned.cell(r, 2).value for r in range(2, 146)]
    original_ids = [base.cell(r, 1).value for r in range(2, 146)]
    original_texts = [base.cell(r, 2).value for r in range(2, 146)]
    if ids != original_ids or len(set(ids)) != 144:
        errors.append({"code": "ID_OR_ORDER_MISMATCH"})
    if texts != original_texts:
        errors.append({"code": "CANDIDATE_TEXT_MISMATCH"})
    if original_ids != [first.cell(r, 1).value for r in range(2, 146)] or original_texts != [first.cell(r, 2).value for r in range(2, 146)]:
        errors.append({"code": "OWN_PHASE1_ID_ORDER_TEXT_NOT_EXACT"})
    if ordered_sha(original_ids) != manifest["ordered_ids_sha256"] or ordered_sha(original_texts) != manifest["candidate_texts_in_order_sha256"]:
        errors.append({"code": "ISSUED_ORDER_MANIFEST_MISMATCH"})
    diffs = readonly_differences(values(issued), values(raw))
    for diff in diffs:
        if diff["sheet"] == SHEET and 2 <= diff["row"] <= 145:
            diff["blind_id"] = returned.cell(diff["row"], 1).value
    answers: list[dict[str, Any]] = []
    flags: list[dict[str, Any]] = []
    leads: list[dict[str, Any]] = []
    legal = 0
    for r in range(2, 146):
        row = {field: returned.cell(r, c).value for c, field in enumerate(FIELDS, 11)}
        available = {tag for tag, c in (("E1", 6), ("E2", 10)) if base.cell(r, c).value not in (None, "")}
        local = first.cell(r, 4).value
        checks = row_checks(row, local, available)
        locator = {"blind_id": ids[r-2], "excel_row": r}
        errors.extend({**locator, **e} for e in checks["errors"])
        flags.extend({**locator, "code": code, "raw_values": row, "own_phase1_local": local} for code in checks["consistency_flags"])
        leads.extend({**locator, "code": code, "raw_values": row} for code in checks["review_leads"])
        legal += checks["legal_categorical_cells"]
        answers.append({**locator, **row, "own_phase1_local_internal_conflict": local})
    for path, expected in inputs:
        if sha(path.read_bytes()) != expected:
            raise ValueError("Read-only source changed during inspection")
    strict_pass = strict_receipt_gate(errors, diffs)
    result = {
        "annotator": person, "raw_sha256": lock["sha256"], "raw_bytes": lock["bytes"],
        "input_hashes_verified": True, "source_hashes_unchanged_after_read": True,
        "issued_sha256": manifest["sha256"], "own_phase1_sha256": p1_sha,
        "rows": len(answers), "unique_ids": len(set(ids)), "unique_candidate_texts": len(set(texts)),
        "id_order_exact": ids == original_ids, "candidate_text_exact": texts == original_texts,
        "headers_exact": not any(e["code"] == "HEADERS_NOT_EXACT" for e in errors),
        "legal_categorical_cells": legal, "expected_categorical_cells": 1152,
        "required_reason_nonblank": sum(isinstance(a["phase2_reason"], str) and bool(a["phase2_reason"].strip()) for a in answers),
        "readonly_difference_counts": dict(Counter(d["kind"] for d in diffs)),
        "readonly_change_sheets": dict(Counter(d["sheet"] for d in diffs if d["kind"] != "BLANK_STORAGE_EQUIVALENT")),
        "readonly_changed_annotation_rows": len({d["row"] for d in diffs if d["sheet"] == SHEET and d["kind"] != "BLANK_STORAGE_EQUIVALENT"}),
        "dropdown_semantic_coverage_exact": validation_signature(base) == validation_signature(returned),
        "zip_safety": safety, "structural_errors": errors, "strict_readonly_receipt_pass": strict_pass,
        "crossfield_flag_count": len(flags), "crossfield_row_count": len({f["blind_id"] for f in flags}),
        "crossfield_codes": dict(Counter(f["code"] for f in flags)),
        "review_lead_count": len(leads), "review_lead_codes": dict(Counter(f["code"] for f in leads)),
        "semantic_validation": "DIAGNOSTIC_ONLY_NO_ANSWER_OVERRIDES", "comparison_not_executed": True,
        "owner_whitespace_exception": "NOT_REQUESTED_YET_NOT_APPROVED",
        "enum_frequencies": {f: dict(Counter(a[f] for a in answers)) for f in FIELDS[:8]},
    }
    new_json(ROOT / f"{suffix}_VALIDATION_V1.json", result)
    new_json(ROOT / f"{suffix}_READONLY_CELL_DIFFERENCES_V1.json", diffs)
    new_json(ROOT / f"{suffix}_ANSWERS_DIAGNOSTIC_ONLY_V1.json", answers)
    new_json(ROOT / f"{suffix}_CROSSFIELD_FLAGS_V1.json", flags)
    new_json(ROOT / f"{suffix}_REVIEW_LEADS_V1.json", leads)
    issued.close()
    raw.close()
    p1.close()
    return result


def main() -> None:
    locks = json.loads((ROOT / "DUAL_PHASE2_RAW_LOCK_MANIFEST_V1.json").read_text(encoding="utf-8"))
    results = [validate_person(lock) for lock in locks]
    allowed = comparison_allowed(*(r["strict_readonly_receipt_pass"] for r in results))
    # No comparison implementation on purpose: a later separately gated task
    # can use an explicitly approved presentation-equivalence overlay.
    summary = {"task": "P1-CORE144-D1-HUMAN-DUAL-PHASE2-RAW-LOCK-AND-COMPARISON-01",
               "raw_lock_complete": True, "isolation": "OWNER_ATTESTED_NOT_MACHINE_PROVEN",
               "individual_validation": results, "ab_comparison_gate_eligible": allowed,
               "ab_comparison_executed": False, "ab_consensus_created": False,
               "expected_or_construction_mapping_loaded": False, "gt_created": False,
               "status": "HUMAN_DECISION_REQUIRED" if not allowed else "RECEIPT_VALIDATED",
               "auto_continue": False}
    new_json(ROOT / "DUAL_PHASE2_RECEIPT_GATE_V1.json", summary)
    for result in results:
        print(json.dumps(result, ensure_ascii=False))
    print("AB_COMPARISON_EXECUTED=False; AUTO_CONTINUE=False")


if __name__ == "__main__":
    main()
