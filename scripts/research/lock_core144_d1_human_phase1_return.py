"""Lock and validate one Core144 D1 human Phase1 workbook return.

The returned workbook is copied byte-for-byte. This script never saves an
``openpyxl`` workbook and therefore cannot normalize or rewrite the raw XLSX.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import stat
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]


LOCKED_NAME = "PAPER1_CORE144_D1_HUMAN_A01_PHASE1_RAW_RETURN_V1.xlsx"
LOCK_MANIFEST_NAME = "PAPER1_CORE144_D1_HUMAN_A01_PHASE1_RAW_LOCK_MANIFEST_V1.json"
VALIDATION_NAME = "PAPER1_CORE144_D1_HUMAN_A01_PHASE1_VALIDATION_V1.json"
PHASE1_SHEETS = [
    "【开始前必读】",
    "【标注表】",
    "【字段说明】",
    "【高危歧义与易错点】",
    "【判断流程】",
    "【虚构示例】",
    "【提交前自检】",
]
PHASE2_SHEETS = [
    "【开始前必读】",
    "【标注表】",
    "【字段说明】",
    "【高危歧义与易错点】",
    "【版本与历史判断】",
    "【权威与机关判断】",
    "【证据判断流程】",
    "【虚构示例】",
    "【冻结证据全文】",
    "【提交前自检】",
]
PHASE1_ENUMS = {
    "text_naturalness": {"NATURAL", "MINOR_ISSUE", "UNNATURAL"},
    "local_internal_conflict": {"YES", "NO", "UNCERTAIN"},
    "self_containment": {"PASS", "FLAG", "UNCERTAIN"},
    "ambiguous_referent": {"YES", "NO", "UNCERTAIN"},
    "meta_or_template_language": {"YES", "NO", "UNCERTAIN"},
}
DEFAULTS = {
    "text_naturalness": "NATURAL",
    "local_internal_conflict": "NO",
    "self_containment": "PASS",
    "ambiguous_referent": "NO",
    "meta_or_template_language": "NO",
}
PHASE1_HEADERS = [
    "blind_id",
    "candidate_text",
    "text_naturalness",
    "local_internal_conflict",
    "self_containment",
    "ambiguous_referent",
    "meta_or_template_language",
    "issue_note",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def stable_hash(values: list[str]) -> str:
    return hashlib.sha256("\n".join(values).encode("utf-8")).hexdigest()


def read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError(f"Expected JSON object: {path}")
    return value


def zip_safety(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
    return {
        "macro_parts": sorted(name for name in names if name.endswith("vbaProject.bin")),
        "external_link_parts": sorted(name for name in names if name.startswith("xl/externalLinks/")),
    }


def validate_phase1(raw_path: Path, original_path: Path, manifest_path: Path) -> dict[str, Any]:
    manifest = read_json(manifest_path)
    if sha256(original_path) != manifest["sha256"] or original_path.stat().st_size != manifest["bytes"]:
        raise AssertionError("Original A01 Phase1 distribution workbook does not match its manifest")

    raw = openpyxl.load_workbook(raw_path, data_only=False, read_only=False)
    original = openpyxl.load_workbook(original_path, data_only=False, read_only=False)
    if raw.sheetnames != PHASE1_SHEETS or original.sheetnames != PHASE1_SHEETS:
        raise AssertionError((raw.sheetnames, original.sheetnames))
    if any(raw[name].sheet_state != "visible" for name in PHASE1_SHEETS):
        raise AssertionError("Raw return contains a hidden Phase1 worksheet")

    sheet = raw["【标注表】"]
    source_sheet = original["【标注表】"]
    headers = [sheet.cell(1, col).value for col in range(1, 9)]
    canonical_headers = [
        value.splitlines()[-1] if isinstance(value, str) else value
        for value in headers
    ]
    if canonical_headers != PHASE1_HEADERS:
        raise AssertionError(headers)
    if sheet.max_row != 145 or sheet.max_column != 8:
        raise AssertionError((sheet.max_row, sheet.max_column))
    if [source_sheet.cell(1, col).value for col in range(1, 9)] != headers:
        raise AssertionError("Phase1 header drift")

    ids: list[str] = []
    texts: list[str] = []
    value_counts: dict[str, Counter[str]] = {
        field: Counter() for field in PHASE1_ENUMS
    }
    nondefault_rows: list[str] = []
    nonblank_note_rows: list[str] = []
    answer_formulas: list[str] = []
    for row in range(2, 146):
        blind_id = sheet.cell(row, 1).value
        text = sheet.cell(row, 2).value
        if blind_id != source_sheet.cell(row, 1).value or text != source_sheet.cell(row, 2).value:
            raise AssertionError(f"ID/candidate drift at Excel row {row}")
        if not isinstance(blind_id, str) or not isinstance(text, str):
            raise AssertionError(f"Missing ID/candidate at Excel row {row}")
        ids.append(blind_id)
        texts.append(text)
        row_values: dict[str, str] = {}
        for col, field in enumerate(PHASE1_ENUMS, start=3):
            value = sheet.cell(row, col).value
            if value not in PHASE1_ENUMS[field]:
                raise AssertionError((row, field, value))
            row_values[field] = value
            value_counts[field][value] += 1
            if isinstance(value, str) and value.startswith("="):
                answer_formulas.append(f"{sheet.cell(row, col).coordinate}:{value}")
        note = sheet.cell(row, 8).value
        is_nondefault = any(row_values[field] != DEFAULTS[field] for field in DEFAULTS)
        if is_nondefault:
            nondefault_rows.append(blind_id)
            if not isinstance(note, str) or not note.strip():
                raise AssertionError(f"Required issue_note missing for {blind_id}")
        if isinstance(note, str) and note.strip():
            nonblank_note_rows.append(blind_id)
        if isinstance(note, str) and note.startswith("="):
            answer_formulas.append(f"H{row}:{note}")

    if len(ids) != 144 or len(set(ids)) != 144:
        raise AssertionError("Phase1 return does not contain 144 unique A01 IDs")
    if stable_hash(ids) != manifest["ordered_ids_sha256"]:
        raise AssertionError("Phase1 A01 ID/order hash mismatch")
    if stable_hash(texts) != manifest["candidate_texts_in_order_sha256"]:
        raise AssertionError("Phase1 A01 candidate text/order hash mismatch")
    if answer_formulas:
        raise AssertionError(answer_formulas)

    # Values/formulas outside the six answer columns must remain identical.
    for name in PHASE1_SHEETS:
        raw_sheet = raw[name]
        old_sheet = original[name]
        if raw_sheet.max_row != old_sheet.max_row or raw_sheet.max_column != old_sheet.max_column:
            raise AssertionError(f"Worksheet dimension drift: {name}")
        for row in range(1, raw_sheet.max_row + 1):
            for col in range(1, raw_sheet.max_column + 1):
                if name == "【标注表】" and row >= 2 and 3 <= col <= 8:
                    continue
                if raw_sheet.cell(row, col).value != old_sheet.cell(row, col).value:
                    raise AssertionError(f"Non-answer cell changed: {name}!{raw_sheet.cell(row, col).coordinate}")

    safety = zip_safety(raw_path)
    if safety["macro_parts"] or safety["external_link_parts"]:
        raise AssertionError(safety)
    return {
        "status": "PASS",
        "records": 144,
        "unique_ids": 144,
        "exact_columns": 8,
        "ordered_ids_sha256": stable_hash(ids),
        "candidate_texts_in_order_sha256": stable_hash(texts),
        "enum_validation": "PASS",
        "required_issue_note_validation": "PASS",
        "nondefault_rows": len(nondefault_rows),
        "nonblank_issue_notes": len(nonblank_note_rows),
        "value_counts": {field: dict(counts) for field, counts in value_counts.items()},
        "answer_formulas": 0,
        "hidden_sheets": 0,
        "macro_parts": 0,
        "external_link_parts": 0,
        "non_answer_content_parity": "PASS",
        "raw_sha256": sha256(raw_path),
        "raw_bytes": raw_path.stat().st_size,
    }


def validate_phase2(phase2_path: Path, manifest_path: Path, phase1_result: dict[str, Any]) -> dict[str, Any]:
    manifest = read_json(manifest_path)
    actual_sha = sha256(phase2_path)
    if actual_sha != manifest["sha256"] or phase2_path.stat().st_size != manifest["bytes"]:
        raise AssertionError("A01 Phase2 workbook does not match its sealed manifest")

    book = openpyxl.load_workbook(phase2_path, data_only=False, read_only=False)
    if book.sheetnames != PHASE2_SHEETS or any(book[name].sheet_state != "visible" for name in PHASE2_SHEETS):
        raise AssertionError(book.sheetnames)
    sheet = book["【标注表】"]
    if sheet.max_row != 145 or sheet.max_column != 20:
        raise AssertionError((sheet.max_row, sheet.max_column))
    ids = [sheet.cell(row, 1).value for row in range(2, 146)]
    texts = [sheet.cell(row, 2).value for row in range(2, 146)]
    if stable_hash(ids) != phase1_result["ordered_ids_sha256"]:
        raise AssertionError("Phase2 ID/order is not identical to locked A01 Phase1")
    if stable_hash(texts) != phase1_result["candidate_texts_in_order_sha256"]:
        raise AssertionError("Phase2 candidate text/order is not identical to locked A01 Phase1")
    for row in range(2, 146):
        for col in range(11, 21):
            if sheet.cell(row, col).value not in (None, ""):
                raise AssertionError(f"Phase2 answer cell is not blank: {sheet.cell(row, col).coordinate}")
    links = 0
    for row in range(2, 146):
        for col in (5, 9):
            value = sheet.cell(row, col).value
            if isinstance(value, str) and value.startswith('=HYPERLINK("https://'):
                links += 1
    forbidden_hits: list[str] = []
    for tab in book:
        for row in tab.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                if re.search(r"\bD1BR-[A-F0-9]{12}\b", value):
                    forbidden_hits.append(f"{tab.title}!{cell.coordinate}:source-id")
                lowered = value.lower()
                if "source_blind_id" in lowered or "target_stealth" in lowered or "ground_truth" in lowered:
                    forbidden_hits.append(f"{tab.title}!{cell.coordinate}:private-field")
    if links != manifest["clickable_official_urls"] or forbidden_hits:
        raise AssertionError((links, forbidden_hits[:10]))
    safety = zip_safety(phase2_path)
    if safety["macro_parts"] or safety["external_link_parts"]:
        raise AssertionError(safety)
    return {
        "status": "PASS",
        "sha256": actual_sha,
        "bytes": phase2_path.stat().st_size,
        "records": 144,
        "ordered_ids_and_candidate_text_parity_with_locked_phase1": "PASS",
        "answer_cells_blank": 144 * 10,
        "clickable_official_urls": links,
        "hidden_sheets": 0,
        "private_mapping_or_label_leakage": 0,
        "macro_parts": 0,
        "external_link_parts": 0,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-raw", type=Path, required=True)
    parser.add_argument("--original-phase1", type=Path, required=True)
    parser.add_argument("--phase1-manifest", type=Path, required=True)
    parser.add_argument("--phase2", type=Path, required=True)
    parser.add_argument("--phase2-manifest", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    for path in (
        args.source_raw,
        args.original_phase1,
        args.phase1_manifest,
        args.phase2,
        args.phase2_manifest,
    ):
        if not path.is_file():
            raise FileNotFoundError(path)

    phase1_result = validate_phase1(args.source_raw, args.original_phase1, args.phase1_manifest)
    phase2_result = validate_phase2(args.phase2, args.phase2_manifest, phase1_result)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    locked_path = args.output_dir / LOCKED_NAME
    lock_manifest_path = args.output_dir / LOCK_MANIFEST_NAME
    validation_path = args.output_dir / VALIDATION_NAME
    for path in (locked_path, lock_manifest_path, validation_path):
        if path.exists():
            raise FileExistsError(f"Refusing to overwrite existing lock artifact: {path}")

    shutil.copyfile(args.source_raw, locked_path)
    if locked_path.read_bytes() != args.source_raw.read_bytes():
        raise AssertionError("Byte-for-byte copy verification failed")
    locked_path.chmod(stat.S_IREAD)
    timestamp = datetime.now(timezone.utc).isoformat()
    validation = {
        "status": "PASS",
        "validated_at_utc": timestamp,
        "phase1": phase1_result,
        "phase2_release_file": phase2_result,
        "release_scope": "HUMAN-A01_ONLY",
        "human_b01_phase2_state": "WITHHELD_PENDING_OWN_PHASE1_RAW_LOCK_AND_VALIDATION",
        "ground_truth_created": False,
    }
    lock_manifest = {
        "artifact": LOCKED_NAME,
        "status": "IMMUTABLE_RAW_LOCKED",
        "locked_at_utc": timestamp,
        "copy_method": "BYTE_FOR_BYTE_NO_EXCEL_RESAVE",
        "source_file_name": args.source_raw.name,
        "source_sha256": phase1_result["raw_sha256"],
        "source_bytes": phase1_result["raw_bytes"],
        "locked_sha256": sha256(locked_path),
        "locked_bytes": locked_path.stat().st_size,
        "locked_file_read_only": bool(
            getattr(locked_path.stat(), "st_file_attributes", 0)
            & getattr(stat, "FILE_ATTRIBUTE_READONLY", 1)
        ),
        "validation_status": "PASS",
        "phase2_release_authorization": "HUMAN-A01_ONLY",
        "phase2_sha256": phase2_result["sha256"],
        "phase2_bytes": phase2_result["bytes"],
    }
    lock_manifest_path.write_text(json.dumps(lock_manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    validation_path.write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "status": "PASS",
                "locked_path": str(locked_path),
                "lock_manifest_path": str(lock_manifest_path),
                "validation_path": str(validation_path),
                "lock_manifest": lock_manifest,
                "validation": validation,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
